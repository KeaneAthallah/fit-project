"""Tanggal pencatatan -- the IDX listing register.

`tanggal pencatatan.xlsx` at the project root lists, per company,
the date its shares were recorded on the exchange: kode, nama,
tanggal. The results grid joins it onto every company-year row so
a reader can see how long a company has been listed -- and filter
on it, because a figure reported for a year before the company was
listed cannot be read at face value.

The register is reference data, not pipeline output, so it is read
straight from the workbook rather than copied into the database:
one small file, parsed once per process and re-read when its mtime
changes.
"""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.core.config import PROJECT_ROOT

PENCATATAN_FILENAME = "tanggal pencatatan.xlsx"

# Indonesian month abbreviations as the IDX prints them, plus the
# English forms for a workbook edited elsewhere. "Agt" is the
# official abbreviation for August; "Agu" is the common variant.
MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "mei": 5, "jun": 6,
    "jul": 7, "agu": 8, "agt": 8, "sep": 9, "okt": 10, "nov": 11, "des": 12,
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5,
    "june": 6, "july": 7, "august": 8, "september": 9, "october": 10,
    "november": 11, "december": 12,
}

# "09 Des 1997", the register's own format: day, month name, year.
_DATE_TEXT_RE = re.compile(r"(\d{1,2})\s+([A-Za-z]+)\.?\s+(\d{4})")

# Legal-form words that carry no identity: "PT Astra Agro Lestari
# Tbk." and "Astra Agro Lestari Tbk" must land on the same key.
_LEGAL_WORDS = {"pt", "tbk", "persero"}


def _normalize(text: str) -> str:
    words = re.sub(r"[^a-z0-9\s]", " ", text.lower()).split()
    return " ".join(w for w in words if w not in _LEGAL_WORDS)


def _parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    m = _DATE_TEXT_RE.match(text)
    if m:
        day, month_name, year = m.groups()
        month = MONTHS.get(month_name.lower())
        if month is not None:
            try:
                return date(int(year), month, int(day))
            except ValueError:
                return None
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


class Pencatatan:
    """The register, indexed both ways.

    The database stores a company as "<KODE> <Nama>", the register
    as two columns, so the join key is the two glued back together.
    The name-only index catches a stored company that lost its
    ticker along the way; the lookup tries the full key first, then
    the name with a leading ticker-like word dropped.
    """

    def __init__(self) -> None:
        self.full: dict[str, date] = {}
        self.by_name: dict[str, date] = {}

    def add(self, kode: str, nama: str, listed: date) -> None:
        name_key = _normalize(nama)
        if not name_key:
            return
        self.full[_normalize(f"{kode} {nama}")] = listed
        self.by_name.setdefault(name_key, listed)

    def get(self, company: str | None) -> date | None:
        if not company:
            return None
        key = _normalize(company)
        if key in self.full:
            return self.full[key]
        if key in self.by_name:
            return self.by_name[key]
        # A stored company that kept its ticker: the full key did
        # not match, so try the name on its own.
        first, _, rest = company.strip().partition(" ")
        if first.isupper() and 1 < len(first) <= 6:
            return self.by_name.get(_normalize(rest))
        return None


_cache: dict[str, tuple[float, Pencatatan]] = {}


def load_pencatatan(path: Path | None = None) -> Pencatatan:
    """The register, re-read only when the file changed.

    Raises when the workbook is missing or holds no usable row: a
    silently empty register would make the "listed before" filter
    hide every company, which reads as a broken table rather than
    as a missing file.
    """
    path = (path or PROJECT_ROOT / PENCATATAN_FILENAME).resolve()
    stamp = path.stat().st_mtime  # raises FileNotFoundError
    cached = _cache.get(str(path))
    if cached is not None and cached[0] == stamp:
        return cached[1]

    register = Pencatatan()
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        for row in ws.iter_rows(values_only=True):
            cells = list(row)
            # The date column anchors the row: the two cells to
            # its left are the code and the name, whatever columns
            # precede them -- the register opens with a "No"
            # column, a variant might open with nothing at all.
            # A header row holds no parseable date and is skipped
            # by the same rule.
            for i, cell in enumerate(cells):
                if i < 2:
                    continue
                listed = _parse_date(cell)
                if listed is None:
                    continue
                kode, nama = cells[i - 2], cells[i - 1]
                if kode is None or nama is None:
                    break
                register.add(str(kode), str(nama), listed)
                break
    finally:
        wb.close()

    if not register.full:
        raise ValueError(f"No tanggal pencatatan could be read from {path}")
    _cache[str(path)] = (stamp, register)
    return register


__all__ = ["Pencatatan", "load_pencatatan", "PENCATATAN_FILENAME"]
