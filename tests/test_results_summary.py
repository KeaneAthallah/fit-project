"""GET /api/results/summary -- the results grid.

One row per company-year, one column per headline figure. The contract that
matters for correctness:

* Rows are the *distinct* (company, year) pairs that have data. A company with
  three years is three rows, and a company with no figures in the requested
  field set is not a row of empty cells.
* A figure absent from a report is `null`, which the grid renders as an em
  dash. That is deliberately different from a figure that extracted as 0: "the
  report has no treasury-share line" and "it holds none" are different claims,
  and only one of them is supported by the source.
* When several values exist for the same (company, year, field) -- which happens
  because balance-sheet figures appear in both the highlights page and the
  statement -- exactly one wins, and the choice is deterministic:
  a hand correction first, then the extractor's own confidence, then the
  lowest id as a stable tie-break.
"""
import io
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.core.config import get_config, load_config, set_config
from app.dashboard.pencatatan import Pencatatan
from app.dashboard.server import create_app
from app.storage.database import get_engine, get_session_factory
from app.storage.repository import Repository


def _value(repo, doc, field, amount, *, year=2024, page=5, method="table", confidence=0.9,
           raw_label=None, currency=None):
    return {
        "company": doc.company,
        "year": year,
        "statement": "balance_sheet",
        "field": field,
        "raw_label": raw_label or field.replace("_", " ").title(),
        "raw_value": None if amount is None else str(amount),
        "normalized_value": amount,
        "unit": None,
        "currency": currency,
        "page": page,
        "extraction_method": method,
        "confidence": confidence,
        "status": "OK",
    }


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.delenv("AI_PROVIDER", raising=False)

    cfg = load_config()
    cfg.database_url = f"sqlite:///{(tmp_path / 'db' / 'processing.db').as_posix()}"
    set_config(cfg)

    app = create_app(cfg)
    app.dependency_overrides[get_config] = lambda: cfg

    Session = get_session_factory(get_engine(cfg))
    with Session() as s:
        repo = Repository(s)
        doc, _ = repo.upsert_document(
            str(tmp_path / "a.pdf"), "hash-a", "PT AAA Tbk", "a.pdf"
        )
        repo.replace_values(
            doc,
            [
                _value(repo, doc, "total_assets", 1_000_000.0, page=5),
                _value(repo, doc, "total_equity", 600_000.0, page=5),
                _value(repo, doc, "sales_and_revenue", 800_000.0, page=7),
                # The IDX classification, verbatim: not a figure, so it
                # rides on text_value and must never become a column or
                # be coerced to a number.
                {
                    **_value(repo, doc, "sub_sector", None),
                    "text_value": "D2. Food & Beverage",
                },
                # Same company, one year earlier.
                _value(repo, doc, "total_assets", 900_000.0, year=2023, page=5),
                # A USD report, in a field the summary grid does not show. The
                # global facet must offer USD because of this row; a column set
                # without finance_costs must not, because no grid cell can be
                # USD because of it.
                _value(repo, doc, "finance_costs", 42_000.0, page=9, currency="USD"),
            ],
        )
        # A second company, to prove rows are not just documents.
        doc2, _ = repo.upsert_document(
            str(tmp_path / "b.pdf"), "hash-b", "PT BBB Tbk", "b.pdf"
        )
        repo.replace_values(
            doc2,
            [
                # One row whose own figures disagree about the currency.
                _value(repo, doc2, "total_assets", 2_000_000.0),
                _value(repo, doc2, "net_income", 120_000.0, currency="USD"),
                # No currency found in the source: a real third state.
                _value(repo, doc2, "finance_costs", 7_000.0),
                # No year at all: the report is undated.
                _value(repo, doc2, "net_income", 150_000.0, year=None),
            ],
        )
        # A company with data but no total_assets must still not appear as an
        # empty row when only total_assets is requested.
        doc3, _ = repo.upsert_document(
            str(tmp_path / "c.pdf"), "hash-c", "PT CCC Tbk", "c.pdf"
        )
        repo.replace_values(
            doc3,
            [
                _value(repo, doc3, "sales", 300_000.0),
                # A second classification, so a reader can tick more
                # than one sector at once.
                {
                    **_value(repo, doc3, "sub_sector", None),
                    "text_value": "11. Crops",
                },
            ],
        )

        # Registered but never read: no extracted values at all. This is the
        # state most of a fresh scan is in, and the reason a company can be
        # missing from the grid without being absent from the filings.
        doc4, _ = repo.upsert_document(
            str(tmp_path / "d.pdf"), "hash-d", "PT DDD Tbk", "d.pdf"
        )
        repo.set_status(doc4, "DISCOVERED")
        doc5, _ = repo.upsert_document(
            str(tmp_path / "e.pdf"), "hash-e", "PT EEE Tbk", "e.pdf"
        )
        repo.set_status(doc5, "DISCOVERED")
        doc6, _ = repo.upsert_document(
            str(tmp_path / "f.pdf"), "hash-f", "PT EEE Tbk", "f.pdf"
        )
        repo.set_status(doc6, "OCR")

        # A company-year whose own arithmetic does not reconcile. The pipeline
        # already recorded this; the grid has to say so.
        repo.replace_validations(
            doc,
            [
                {
                    "company": "PT AAA Tbk",
                    "year": 2024,
                    "check_name": "assets_equals_liabilities_plus_equity",
                    "status": "ERROR",
                    "severity": "ERROR",
                    "expected": "1600000.0",
                    "actual": "1000000.0",
                }
            ],
        )

    with TestClient(app) as c:
        yield c

    set_config(load_config())


def _rows(client, **params):
    r = client.get("/api/results/summary", params=params)
    assert r.status_code == 200
    return r.json()


# --------------------------------------------------------------------------
# shape


def test_one_row_per_company_year(client):
    body = _rows(client)
    keys = {(row["company"], row["year"]) for row in body["items"]}
    assert keys == {
        ("PT AAA Tbk", 2024),
        ("PT AAA Tbk", 2023),
        ("PT BBB Tbk", 2024),
        ("PT BBB Tbk", None),
        ("PT CCC Tbk", 2024),
    }
    assert body["pagination"]["total"] == 5


def test_default_columns_are_the_headline_figures_in_reading_order(client):
    """The grid opens on the figures a reader works down, and nothing else.

    A fixed list is the point: the same eight columns for every company, so two
    rows are comparable at a glance. A headline figure no report in view
    contained is dropped rather than shown as a column of em dashes, and a
    figure outside the list is reached through ?fields= or the exports.
    """
    body = _rows(client)
    assert body["fields"] == [
        "total_assets",
        "total_equity",
        "sales_and_revenue",
        "sales",
        "net_income",
    ]

    # Labels travel with the response so header text has one source of truth.
    assert body["labels"]["total_assets"] == "Total assets"
    assert body["labels"]["net_income"] == "Profit for the year"


def test_a_field_outside_the_headline_list_is_not_a_default_column(client):
    """finance_costs was extracted and is real. It is still not the summary."""
    body = _rows(client)
    assert "finance_costs" not in body["fields"]
    # It is one `?fields=` away, not gone.
    assert _rows(client, fields="finance_costs")["fields"] == ["finance_costs"]


def test_a_field_that_was_never_extracted_gets_no_column(client):
    """A label in the mapping table is not evidence that a report had it."""
    body = _rows(client)
    assert "treasury_shares_carrying_value" not in body["fields"]


def test_columns_can_still_be_widened_past_the_headline_figures(client):
    """The default list is a default, not a ceiling."""
    extra = "gross_profit,finance_costs"
    body = _rows(client, fields=extra)
    assert body["fields"] == extra.split(",")


def test_custom_fields_are_returned_in_the_requested_order(client):
    body = _rows(client, fields="net_income,total_assets")
    assert body["fields"] == ["net_income", "total_assets"]


def test_blank_fields_falls_back_to_the_defaults(client):
    """A stray `?fields=` should not be a hard error; it means "no preference"."""
    assert _rows(client, fields="")["fields"] == _rows(client)["fields"]
    assert _rows(client, fields=" , ")["fields"] == _rows(client)["fields"]


def test_unknown_field_is_rejected_rather_than_silently_dropped(client):
    """Returning a narrower table would hide the typo instead of fixing it."""
    r = client.get("/api/results/summary", params={"fields": "total_assets,net_incom"})
    assert r.status_code == 400
    assert "net_incom" in r.json()["detail"]


# --------------------------------------------------------------------------
# values and absence


def test_figures_land_in_the_right_row_and_column(client):
    body = _rows(client)
    aaa_2024 = next(r for r in body["items"] if (r["company"], r["year"]) == ("PT AAA Tbk", 2024))
    assert aaa_2024["cells"]["total_assets"]["normalized_value"] == 1_000_000.0
    assert aaa_2024["cells"]["sales_and_revenue"]["normalized_value"] == 800_000.0
    assert aaa_2024["cells"]["net_income"] is None


def test_absent_figure_is_null_not_zero(client):
    """finance_costs exists in the data, but not for PT CCC.

    Reporting 0 for CCC would assert the company had no finance costs, which is
    a different and unsupported claim from "this line was not extracted".
    """
    body = _rows(client, fields="finance_costs,sales")
    ccc = next(r for r in body["items"] if r["company"] == "PT CCC Tbk")
    aaa = next(r for r in body["items"] if (r["company"], r["year"]) == ("PT AAA Tbk", 2024))
    assert ccc["cells"]["finance_costs"] is None
    assert ccc["cells"]["sales"]["normalized_value"] == 300_000.0
    assert aaa["cells"]["finance_costs"]["normalized_value"] == 42_000.0


def test_sales_and_revenue_are_separate_columns(client):
    body = _rows(client)
    aaa_2024 = next(r for r in body["items"] if (r["company"], r["year"]) == ("PT AAA Tbk", 2024))
    ccc = next(r for r in body["items"] if r["company"] == "PT CCC Tbk")
    # A report that says "Penjualan dan pendapatan usaha" populates
    # sales_and_revenue and leaves sales empty. Collapsing the two would report
    # a sales figure as the combined line.
    assert aaa_2024["cells"]["sales"] is None
    assert aaa_2024["cells"]["sales_and_revenue"]["normalized_value"] == 800_000.0
    assert ccc["cells"]["sales"]["normalized_value"] == 300_000.0
    assert ccc["cells"]["sales_and_revenue"] is None


def test_undated_report_still_gets_a_row(client):
    body = _rows(client)
    undated = [r for r in body["items"] if r["company"] == "PT BBB Tbk" and r["year"] is None]
    assert len(undated) == 1
    assert undated[0]["cells"]["net_income"]["normalized_value"] == 150_000.0


# --------------------------------------------------------------------------
# picking one value when duplicates exist


def test_higher_confidence_wins_between_duplicate_figures(client):
    """The same figure can be extracted from two documents.

    One document per (company, year, field) is unique, so duplicates arise
    across documents -- a company whose highlights PDF and full report were
    both processed. Both extracts are kept (the source page differs), so the
    grid has to choose. Confidence is the extractor's own signal, so that is
    what breaks the tie.
    """
    Session = get_session_factory(get_engine(get_config()))
    with Session() as s:
        repo = Repository(s)
        # The annual report, indexed with lower confidence than the highlights
        # PDF that is already in the fixture.
        doc, _ = repo.upsert_document(
            "annual.pdf", "hash-a-annual", "PT AAA Tbk", "annual.pdf"
        )
        repo.replace_values(doc, [_value(repo, doc, "total_assets", 1_000_000.0, confidence=0.55)])

    body = _rows(client)
    aaa_2024 = next(r for r in body["items"] if (r["company"], r["year"]) == ("PT AAA Tbk", 2024))
    assert aaa_2024["cells"]["total_assets"]["normalized_value"] == 1_000_000.0
    assert aaa_2024["cells"]["total_assets"]["confidence"] == 0.9
    # Still one row for that company-year, not two.
    assert sum(
        1 for r in body["items"] if (r["company"], r["year"]) == ("PT AAA Tbk", 2024)
    ) == 1


def test_a_hand_correction_outranks_a_higher_confidence_extract(client):
    """The user knows more than the confidence score does.

    If a manual fix lost to a 0.99-confidence extraction, the grid would keep
    showing a number the user has already said is wrong -- which is the whole
    reason the correction was made.
    """
    listed = client.get("/api/values", params={"field": "total_assets", "company": "PT AAA Tbk"})
    value_id = next(
        v["id"] for v in listed.json()["items"] if v["normalized_value"] == 1_000_000.0
    )
    patched = client.patch(f"/api/values/{value_id}", json={"normalized_value": 1_234_567.0})
    assert patched.status_code == 200

    body = _rows(client)
    aaa_2024 = next(r for r in body["items"] if (r["company"], r["year"]) == ("PT AAA Tbk", 2024))
    cell = aaa_2024["cells"]["total_assets"]
    assert cell["normalized_value"] == 1_234_567.0
    assert cell["is_edited"] is True


def test_reverting_puts_the_extractor_figure_back_in_the_grid(client):
    listed = client.get("/api/values", params={"field": "total_assets", "company": "PT AAA Tbk"})
    value_id = listed.json()["items"][0]["id"]
    client.patch(f"/api/values/{value_id}", json={"normalized_value": 7.0})
    client.patch(f"/api/values/{value_id}", json={"revert": True})

    body = _rows(client)
    aaa_2024 = next(r for r in body["items"] if (r["company"], r["year"]) == ("PT AAA Tbk", 2024))
    cell = aaa_2024["cells"]["total_assets"]
    assert cell["normalized_value"] == 1_000_000.0
    assert cell["is_edited"] is False


# --------------------------------------------------------------------------
# filters and paging


def test_filter_by_company(client):
    body = _rows(client, company="PT BBB Tbk")
    assert {r["company"] for r in body["items"]} == {"PT BBB Tbk"}


def test_filter_by_year(client):
    body = _rows(client, year=2023)
    assert {(r["company"], r["year"]) for r in body["items"]} == {("PT AAA Tbk", 2023)}


def test_paging_splits_rows_without_repeating_any(client):
    first = _rows(client, page=1, page_size=3)
    second = _rows(client, page=2, page_size=3)
    assert len(first["items"]) == 3
    assert len(second["items"]) == 2
    assert first["pagination"]["pages"] == 2
    assert first["pagination"]["page_size"] == 3

    keys = {(r["company"], r["year"]) for r in first["items"] + second["items"]}
    assert len(keys) == 5


def test_page_beyond_the_end_is_empty_not_an_error(client):
    body = _rows(client, page=99, page_size=10)
    assert body["items"] == []
    assert body["pagination"]["total"] == 5


# --------------------------------------------------------------------------
# currency
#
# The option list is part of the feature: a dropdown that offers a currency
# which then matches nothing reads as a broken filter, not as an empty result.


def test_currency_choices_exclude_currencies_that_cannot_match_a_cell(client):
    """finance_costs is USD here, and it is not a summary column.

    The global facet lists USD because the value exists. Offering USD on the
    grid would return zero rows and look like the filter was ignored, so the
    summary derives its own list from the fields it actually shows.
    """
    global_codes = {c["currency"] for c in client.get("/api/values/facets").json()["currencies"]}
    assert "USD" in global_codes, "fixture should make USD reachable globally"

    # A column set with no USD in it must not offer USD.
    assert {c["currency"] for c in _rows(client, fields="total_assets")["currencies"]} == {"none"}

    # Asking for the field that really is USD brings it back, proving the list
    # is scoped by field rather than simply blind to USD.
    assert {c["currency"] for c in _rows(client, fields="finance_costs")["currencies"]} == {
        "USD",
        "none",
    }


def test_choosing_a_currency_does_not_shrink_the_choice_list(client):
    """The list ignores its own filter, or the filter could only be set once."""
    assert _rows(client, fields="finance_costs", currency="USD")["currencies"] == _rows(
        client, fields="finance_costs"
    )["currencies"]


def test_the_currency_in_force_stays_offered_when_the_company_has_none(client):
    """A filter that cannot be seen cannot be cleared.

    The choice list is scoped to the company and year on screen, so a company
    with no USD figure drops USD from the dropdown. If the select is left
    holding `currency=USD` with no matching option, the browser renders the
    first option instead: the control reads "All currencies" while the query
    still filters on USD, and the reader has no way back. The currency in force
    is therefore always offered, and a currency that matches nothing reports a
    count of zero rather than vanishing.
    """
    body = _rows(client, company="PT BBB Tbk", currency="USD", fields="finance_costs")
    assert body["items"] == [], "fixture company holds no USD figure"
    offered = {c["currency"]: c["count"] for c in body["currencies"]}
    assert offered["USD"] == 0, "the applied filter must be listed, at count zero"

    # Same request spelled in a different case is the same filter, so it must
    # not appear as a second, empty option beside the real one.
    lower = _rows(client, company="PT BBB Tbk", currency="usd", fields="finance_costs")
    assert [c["currency"] for c in lower["currencies"]] == [
        c["currency"] for c in body["currencies"]
    ]
    assert len(lower["currencies"]) == len(body["currencies"])

    # An undetected bucket behaves the same way rather than disappearing.
    none_only = _rows(client, company="PT AAA Tbk", currency="none", fields="finance_costs")
    assert none_only["items"] == []
    assert {c["currency"]: c["count"] for c in none_only["currencies"]}["none"] == 0


def test_filtering_by_currency_narrows_the_rows(client):
    usd = _rows(client, fields="finance_costs", currency="USD")["items"]
    assert {r["company"] for r in usd} == {"PT AAA Tbk"}
    assert all(
        cell["currency"] == "USD" for row in usd for cell in row["cells"].values() if cell
    )

    # Undetected is a real bucket, not a synonym for "no rows".
    undetected = _rows(client, fields="finance_costs", currency="none")["items"]
    assert {r["company"] for r in undetected} == {"PT BBB Tbk"}
    assert all(
        cell["currency"] is None for row in undetected for cell in row["cells"].values() if cell
    )


def test_row_says_mixed_when_its_own_figures_disagree_about_currency(client):
    body = _rows(client, company="PT BBB Tbk", year=2024)
    assert {row["currency"] for row in body["items"]} == {"Mixed"}

    # Filtering by a currency still surfaces the row, with only its own figures.
    body = _rows(client, company="PT BBB Tbk", year=2024, currency="USD")
    assert body["items"][0]["cells"]["net_income"]["normalized_value"] == 120_000.0
    assert body["items"][0]["cells"]["total_assets"] is None

    # The undated row of the same company reports its own currency state
    # rather than inheriting the dated row's answer.
    undated = _rows(client, company="PT BBB Tbk")["items"]
    assert [r["year"] for r in undated if r["year"] is None] == [None]


def test_currency_filter_is_case_insensitive(client):
    assert _rows(client, currency="usd", fields="finance_costs")["items"] == _rows(
        client, currency="USD", fields="finance_costs"
    )["items"]


def test_the_subsector_filter_narrows_the_grid(client):
    """The declared classification is a filter like any other, and
    the option list is scoped the way the currencies are: only what
    the other filters can still reach."""
    body = _rows(client, subsector="D2. Food & Beverage")
    assert {r["company"] for r in body["items"]} == {"PT AAA Tbk"}
    assert {s["subsector"] for s in body["subsectors"]} == {
        "D2. Food & Beverage",
        "11. Crops",
        "none",
    }


def test_the_undeclared_subsector_is_its_own_bucket(client):
    """A company that prints no classification on any filing is not the
    empty string, and must not be conflated with a sector literally
    named "none"."""
    body = _rows(client, subsector="none")
    keys = {(r["company"], r["year"]) for r in body["items"]}
    assert ("PT BBB Tbk", 2024) in keys
    assert ("PT BBB Tbk", None) in keys
    # The other two companies declare a classification on some filing,
    # so none of their rows are unclassified.
    assert ("PT AAA Tbk", 2024) not in keys
    assert ("PT CCC Tbk", 2024) not in keys


def test_the_undeclared_bucket_can_be_picked_alongside_real_sectors(client):
    body = _rows(client, subsector=["none", "D2. Food & Beverage"])
    assert {r["company"] for r in body["items"]} == {"PT AAA Tbk", "PT BBB Tbk"}


def test_the_subsector_filter_accepts_several_selections(client):
    """Selections travel as repeated parameters rather than one joined
    string, because a classification can itself contain a comma --
    "43. Textile, Garment" is a real IDX code."""
    body = _rows(client, subsector=["D2. Food & Beverage", "11. Crops"])
    assert {r["company"] for r in body["items"]} == {"PT AAA Tbk", "PT CCC Tbk"}
    # Every selection is still offered, scoped like the currency list.
    assert {s["subsector"] for s in body["subsectors"]} == {
        "D2. Food & Beverage",
        "11. Crops",
        "none",
    }


def test_the_subsector_travels_with_the_row(client):
    body = _rows(client)
    aaa = next(r for r in body["items"] if (r["company"], r["year"]) == ("PT AAA Tbk", 2024))
    assert aaa["subsector"] == "D2. Food & Beverage"
    # Stated as filed, not retitled: the classification is a code.
    bbb = next(r for r in body["items"] if r["company"] == "PT BBB Tbk")
    assert bbb["subsector"] is None
    ccc = next(r for r in body["items"] if r["company"] == "PT CCC Tbk")
    assert ccc["subsector"] == "11. Crops"


def test_the_subsector_belongs_to_the_company_not_the_year(client):
    """A company has one classification, not one per filing year. The
    2023 report never printed it, yet the row still carries it -- a
    sector is not something a company changes between reports."""
    body = _rows(client, year=2023)
    assert len(body["items"]) == 1
    assert body["items"][0]["subsector"] == "D2. Food & Beverage"


def test_the_no_net_loss_filter_keeps_only_profitable_years(client, tmp_path):
    """A year that lost money is dropped, and so is a year with no
    profit figure at all: no loss is not evidence of a profit."""
    Session = get_session_factory(
        get_engine(client.app.dependency_overrides[get_config]())
    )
    with Session() as s:
        repo = Repository(s)
        doc, _ = repo.upsert_document(
            str(tmp_path / "p.pdf"), "hash-p", "PT PROFIT Tbk", "p.pdf"
        )
        repo.replace_values(
            doc,
            [
                _value(repo, doc, "total_profit_loss", 50_000.0, year=2024),
                _value(repo, doc, "total_profit_loss", -10_000.0, year=2023),
            ],
        )

    kept = {
        (r["company"], r["year"]) for r in _rows(client, profitable=True)["items"]
    }
    assert ("PT PROFIT Tbk", 2024) in kept
    assert ("PT PROFIT Tbk", 2023) not in kept
    # The fixture holds no profit figures, so none of it counts as
    # healthy -- a missing number is not a clean year.
    assert ("PT AAA Tbk", 2024) not in kept


def test_the_no_net_loss_filter_narrows_the_export_too(client, tmp_path):
    Session = get_session_factory(
        get_engine(client.app.dependency_overrides[get_config]())
    )
    with Session() as s:
        repo = Repository(s)
        doc, _ = repo.upsert_document(
            str(tmp_path / "p.pdf"), "hash-p", "PT PROFIT Tbk", "p.pdf"
        )
        repo.replace_values(
            doc,
            [
                _value(repo, doc, "total_profit_loss", 50_000.0, year=2024),
                _value(repo, doc, "total_profit_loss", -10_000.0, year=2023),
            ],
        )

    wb = _workbook(_export(client, profitable=True))
    assert [r[0] for r in _sheet_rows(wb["Summary"])] == ["PT PROFIT Tbk"]
    # The filter the file was taken under is recorded in it.
    text = "\n".join(str(c.value) for r in wb["Notes"].iter_rows() for c in r if c.value)
    assert "no net loss" in text


def _register(entries: dict[str, date]) -> Pencatatan:
    """A register seeded from full company names: the first
    word stands in for the ticker and the rest for the
    name, which is how the register and the database split
    one company between them."""
    register = Pencatatan()
    for company, listed in entries.items():
        words = company.split()
        register.add(words[0], " ".join(words[1:]), listed)
    return register


def test_the_rugi_terus_filter_keeps_only_loss_years(client, tmp_path):
    """The mirror of laba terus: only years that reported a
    loss, and a year with no profit figure is not one of
    them -- a missing number is evidence of neither
    direction."""
    Session = get_session_factory(
        get_engine(client.app.dependency_overrides[get_config]())
    )
    with Session() as s:
        repo = Repository(s)
        doc, _ = repo.upsert_document(
            str(tmp_path / "p.pdf"), "hash-p", "PT PROFIT Tbk", "p.pdf"
        )
        repo.replace_values(
            doc,
            [
                _value(repo, doc, "total_profit_loss", 50_000.0, year=2024),
                _value(repo, doc, "total_profit_loss", -10_000.0, year=2023),
            ],
        )

    kept = {
        (r["company"], r["year"]) for r in _rows(client, profitable="rugi")["items"]
    }
    assert ("PT PROFIT Tbk", 2023) in kept
    assert ("PT PROFIT Tbk", 2024) not in kept
    # The fixture holds no profit figures, so none of it
    # counts as a loss either.
    assert ("PT AAA Tbk", 2024) not in kept


def test_the_rugi_filter_narrows_the_export_too(client, tmp_path):
    Session = get_session_factory(
        get_engine(client.app.dependency_overrides[get_config]())
    )
    with Session() as s:
        repo = Repository(s)
        doc, _ = repo.upsert_document(
            str(tmp_path / "p.pdf"), "hash-p", "PT PROFIT Tbk", "p.pdf"
        )
        repo.replace_values(
            doc,
            [
                _value(repo, doc, "total_profit_loss", 50_000.0, year=2024),
                _value(repo, doc, "total_profit_loss", -10_000.0, year=2023),
            ],
        )

    wb = _workbook(_export(client, profitable="rugi"))
    assert [r[0] for r in _sheet_rows(wb["Summary"])] == ["PT PROFIT Tbk"]
    # The filter the file was taken under is recorded in it.
    text = "\n".join(str(c.value) for r in wb["Notes"].iter_rows() for c in r if c.value)
    assert "net loss" in text


def test_the_pencatatan_column_carries_the_listing_date(client, monkeypatch):
    """Every row carries the date its company was listed, so
    a reader can see how long the numbers span. A company
    the register does not know shows null -- never a
    guessed date."""
    monkeypatch.setattr(
        "app.dashboard.api.load_pencatatan",
        lambda: _register({"PT AAA Tbk": date(1997, 12, 9)}),
    )
    body = _rows(client)
    aaa = next(
        r for r in body["items"]
        if r["company"] == "PT AAA Tbk" and r["year"] == 2024
    )
    assert aaa["pencatatan"] == "1997-12-09"
    bbb = next(
        r for r in body["items"]
        if r["company"] == "PT BBB Tbk" and r["year"] == 2024
    )
    assert bbb["pencatatan"] is None


def test_the_listing_date_filter_keeps_only_companies_listed_before(
    client, monkeypatch,
):
    """A company listed in 2022 cannot honestly be part of a
    "before 2020" answer, and neither can one the register
    does not know: an unknown date is not a "yes"."""
    monkeypatch.setattr(
        "app.dashboard.api.load_pencatatan",
        lambda: _register({
            "PT AAA Tbk": date(1997, 12, 9),
            "PT BBB Tbk": date(2022, 8, 4),
        }),
    )
    body = _rows(client, pencatatan_before=2020)
    companies = {r["company"] for r in body["items"]}
    assert companies == {"PT AAA Tbk"}


def test_the_listing_date_filter_narrows_the_export_too(client, monkeypatch):
    monkeypatch.setattr(
        "app.dashboard.api.load_pencatatan",
        lambda: _register({
            "PT AAA Tbk": date(1997, 12, 9),
            "PT BBB Tbk": date(2022, 8, 4),
        }),
    )

    wb = _workbook(_export(client, pencatatan_before=2020))
    assert [r[0] for r in _sheet_rows(wb["Summary"])] == [
        "PT AAA Tbk", "PT AAA Tbk",
    ]
    # The Pencatatan column is the fifth, and carries the
    # date as a real date cell.
    listed = wb["Summary"]["E2"].value
    assert listed is not None
    assert listed.date() == date(1997, 12, 9)
    # The filter the file was taken under is recorded in it.
    text = "\n".join(str(c.value) for r in wb["Notes"].iter_rows() for c in r if c.value)
    assert "listed before 2020" in text


def test_coverage_names_the_companies_that_have_not_been_read(client):
    """The gap must be stated, not left as a silent omission.

    A company with queued PDFs is absent from the grid. Without this the only
    visible fact is "PT DDD is not in the report", which reads as "there is no
    filing" when the truth is "we have not opened it yet".
    """
    body = client.get("/api/results/coverage").json()

    assert body["companies_discovered"] == 5
    assert body["companies_with_values"] == 3
    assert body["companies_without_values"] == 2
    assert body["documents_total"] == 6
    assert body["documents_with_values"] == 3

    missing = {m["company"]: m for m in body["missing"]}
    assert set(missing) == {"PT DDD Tbk", "PT EEE Tbk"}
    assert missing["PT DDD Tbk"]["by_status"] == {"DISCOVERED": 1}
    # Two documents for one company, so the per-status split has to add up.
    assert missing["PT EEE Tbk"]["documents"] == 2
    assert missing["PT EEE Tbk"]["by_status"] == {"DISCOVERED": 1, "OCR": 1}

    # A company that did produce data is never listed as missing, however many
    # of its other documents are still queued.
    assert "PT AAA Tbk" not in missing


def test_a_row_whose_arithmetic_fails_says_so(client):
    """The identities are already run; the grid has to report the verdict.

    A row whose own figures do not add up was previously indistinguishable
    from one that reconciles, which made 'these are the figures' read as
    'these figures are correct'.
    """
    body = _rows(client)
    aaa = next(r for r in body["items"] if (r["company"], r["year"]) == ("PT AAA Tbk", 2024))
    assert aaa["failed_checks"] == {
        "assets_equals_liabilities_plus_equity": "ERROR",
    }

    # A company-year that reconciled must not be decorated with a warning.
    bbb = next(r for r in body["items"] if r["company"] == "PT BBB Tbk" and r["year"] == 2024)
    assert bbb["failed_checks"] == {}

    # NOT_APPLICABLE means "not enough data to judge", which is not a failure.
    ccc = next(r for r in body["items"] if r["company"] == "PT CCC Tbk")
    assert ccc["failed_checks"] == {}


def test_failed_checks_follow_the_company_filter(client):
    body = _rows(client, company="PT AAA Tbk", year=2024)
    assert body["items"][0]["failed_checks"] == {
        "assets_equals_liabilities_plus_equity": "ERROR",
    }
    other = _rows(client, company="PT BBB Tbk")
    assert all(r["failed_checks"] == {} for r in other["items"])


def test_no_coverage_notice_when_every_company_has_data(client, tmp_path, monkeypatch):
    """The notice must disappear once the backlog is cleared, not linger."""
    body = client.get("/api/results/coverage").json()
    assert body["companies_without_values"] == 2

    cfg = client.app.dependency_overrides[get_config]()
    Session = get_session_factory(get_engine(cfg))
    with Session() as s:
        repo = Repository(s)
        for doc in repo.all_documents():
            if not repo.values_for_document(doc):
                doc_q, _ = repo.upsert_document(
                    str(tmp_path / f"{doc.id}-x.pdf"), f"x{doc.id}", doc.company, "x.pdf"
                )
                repo.replace_values(
                    doc_q, [_value(repo, doc_q, "total_assets", 1.0)]
                )

        body = client.get("/api/results/coverage").json()
        assert body["companies_without_values"] == 0
        assert body["missing"] == []


def _seed_conflicting(client, tmp_path):
    """Two reports of the same company that disagree about one figure.

    The printed digits are identical; one document's stray unit scaled them by
    1,000,000, and that inflated copy carries the higher confidence.
    """
    Session = get_session_factory(get_engine(client.app.dependency_overrides[get_config]()))
    with Session() as s:
        repo = Repository(s)
        wrong, _ = repo.upsert_document(
            str(tmp_path / "w.pdf"), "hash-w", "PT DUP Tbk", "w.pdf"
        )
        right, _ = repo.upsert_document(
            str(tmp_path / "r.pdf"), "hash-r", "PT DUP Tbk", "r.pdf"
        )
        wrong_v = _value(repo, wrong, "net_income", 17_527_130_084_000_000.0,
                         confidence=0.855)
        wrong_v["raw_value"] = "17.527.130.084"
        right_v = _value(repo, right, "net_income", 17_527_130_084.0,
                         confidence=0.54)
        right_v["raw_value"] = "17.527.130.084"
        repo.replace_values(wrong, [wrong_v])
        repo.replace_values(right, [right_v])
    return wrong, right


class TestCompanyProfile:
    """The profile answers a different question from the grid: not "how do these
    companies compare" but "what do we hold for this one". It must not invent a
    different answer to the same figures."""

    def test_profile_aggregates_every_filing_for_the_company(self, client):
        body = client.get("/api/companies/PT%20AAA%20Tbk").json()
        assert body["company"] == "PT AAA Tbk"
        assert body["totals"]["documents"] == 1
        # Both years the fixture wrote must appear, newest last for charting.
        assert body["years"] == [2023, 2024]
        assert body["metrics"]["total_assets"][-1]["normalized_value"] == 1_000_000.0

    def test_profile_reports_disputes_the_same_way_the_grid_does(self, client, tmp_path):
        """Two views of one figure that disagree about whether it is settled
        would be worse than either view alone."""
        _seed_conflicting(client, tmp_path)
        cell = _rows(client, company="PT DUP Tbk")["items"][0]["cells"]["net_income"]
        assert cell["disputed"] is True

        profile = client.get("/api/companies/PT%20DUP%20Tbk").json()
        series = profile["metrics"]["net_income"]
        assert series[-1]["disputed"] is True
        assert profile["quality"]["disputed_cells"] >= 1

    def test_missing_years_are_explicit_rather_than_skipped(self, client):
        """A gap in the series has to render as a gap. Dropping the year would
        silently compress the time axis and make a two-year gap look like one."""
        profile = client.get("/api/companies/PT%20AAA%20Tbk").json()
        # Assets exist for both years, so the axis is 2023..2024 for them.
        assert [p["year"] for p in profile["metrics"]["total_assets"]] == [2023, 2024]
        # Equity was only reported for 2024; 2023 must still appear, as a null.
        equity = profile["metrics"]["total_equity"]
        assert [p["year"] for p in equity] == [2023, 2024]
        assert equity[0]["normalized_value"] is None
        assert equity[1]["normalized_value"] == 600_000.0

    def test_unknown_company_is_a_404(self, client):
        assert client.get("/api/companies/PT%20NOPE%20Tbk").status_code == 404

    def test_list_endpoint_still_works_alongside_the_profile_route(self, client):
        """The profile route sits under the same prefix; it must not swallow the
        plain company list."""
        assert "items" in client.get("/api/companies").json()


class TestDisputedFigures:
    """When sources disagree the grid must not pick a winner.

    Confidence is anti-correlated with correctness in this data: a value
    inflated by 1,000,000x scored 0.855 and beat the correct reading at 0.54.
    Ranking by it therefore selects the wrong side rather than merely hiding the
    conflict, so every material disagreement is reported with its candidates.
    """

    @staticmethod
    def _conflicting(client, tmp_path):
        """Kept as a method so the existing tests read unchanged; the seeding
        itself now lives at module level, where the profile tests can share it."""
        return _seed_conflicting(client, tmp_path)

    def test_conflicting_figures_are_reported_not_silently_picked(self, client, tmp_path):
        self._conflicting(client, tmp_path)
        body = _rows(client, company="PT DUP Tbk")
        cell = body["items"][0]["cells"]["net_income"]

        assert cell["disputed"] is True
        # Both readings travel with the cell so a reader can judge.
        amounts = sorted(c["normalized_value"] for c in cell["candidates"])
        assert amounts == [17_527_130_084.0, 17_527_130_084_000_000.0]

    def test_the_inflated_copy_is_never_presented_as_settled(self, client, tmp_path):
        """The top-ranked candidate may still be the wrong one, so the grid has
        to keep saying the figure is unresolved."""
        self._conflicting(client, tmp_path)
        cell = _rows(client, company="PT DUP Tbk")["items"][0]["cells"]["net_income"]
        assert cell["disputed"] is True
        assert cell["candidates"][0]["confidence"] == 0.855

    def test_agreeing_duplicates_are_not_disputes(self, client, tmp_path):
        """Two reports agreeing on the same figure is corroboration, not a
        conflict, and must not decorate the grid."""
        Session = get_session_factory(get_engine(client.app.dependency_overrides[get_config]()))
        with Session() as s:
            repo = Repository(s)
            a, _ = repo.upsert_document(
                str(tmp_path / "a2.pdf"), "hash-a2", "PT SAME Tbk", "a.pdf"
            )
            b, _ = repo.upsert_document(
                str(tmp_path / "b2.pdf"), "hash-b2", "PT SAME Tbk", "b.pdf"
            )
            repo.replace_values(a, [_value(repo, a, "net_income", 5_000_000.0)])
            repo.replace_values(b, [_value(repo, b, "net_income", 5_000_000.0)])
        cell = _rows(client, company="PT SAME Tbk")["items"][0]["cells"]["net_income"]
        assert "disputed" not in cell

    def _two_readings(self, client, tmp_path, low: float, high: float) -> dict:
        """Two reports of one company reading the same figure two ways."""
        Session = get_session_factory(get_engine(client.app.dependency_overrides[get_config]()))
        with Session() as s:
            repo = Repository(s)
            a, _ = repo.upsert_document(
                str(tmp_path / "lo.pdf"), "hash-lo", "PT GAP Tbk", "lo.pdf"
            )
            b, _ = repo.upsert_document(
                str(tmp_path / "hi.pdf"), "hash-hi", "PT GAP Tbk", "hi.pdf"
            )
            repo.replace_values(a, [_value(repo, a, "net_income", low)])
            repo.replace_values(b, [_value(repo, b, "net_income", high)])
        return _rows(client, company="PT GAP Tbk")["items"][0]["cells"]["net_income"]

    @pytest.mark.parametrize(
        "low,high,expected",
        [
            # 29% of the larger reading: rounding-scale noise.
            (1_000_000_000.0, 1_400_000_000.0, False),
            # A small reading beside a large one is judged on its own size, not
            # measured against the small one and blown up.
            (100.0, 149.0, False),
            # Exactly half the larger reading. The rule is "more than 50%", so
            # the boundary itself is not a dispute.
            (1_000_000_000.0, 2_000_000_000.0, False),
            # Just past it.
            (1_000_000_000.0, 2_000_001_000.0, True),
            (100.0, 201.0, True),
            (1_000_000.0, 1_800_000_000.0, True),
        ],
    )
    def test_the_threshold_is_half_the_larger_reading(self, client, tmp_path, low, high, expected):
        cell = self._two_readings(client, tmp_path, low, high)
        assert ("disputed" in cell) is expected, (
            f"{low} vs {high} is {(high - low) / high:.1%} of the larger reading "
            f"and should {'be disputed' if expected else 'be treated as agreement'}"
        )

    def test_a_conflict_reports_both_readings(self, client, tmp_path):
        """Disputed or not, disagreeing readings must be retrievable, otherwise
        a reader cannot check the judgement."""
        cell = self._two_readings(client, tmp_path, 1_000_000_000.0, 2_100_000_000.0)
        assert cell["disputed"] is True
        assert sorted(c["normalized_value"] for c in cell["candidates"]) == [
            1_000_000_000.0,
            2_100_000_000.0,
        ]

    def test_a_hand_correction_settles_the_dispute(self, client, tmp_path):
        """Choosing a reading is recorded as a correction, and a correction is
        authoritative: the conflict stops being reported."""
        wrong, _ = self._conflicting(client, tmp_path)
        cell = _rows(client, company="PT DUP Tbk")["items"][0]["cells"]["net_income"]
        assert cell["disputed"] is True

        resp = client.patch(
            f"/api/values/{cell['id']}",
            json={"normalized_value": 17_527_130_084.0},
        )
        assert resp.status_code == 200

        cell = _rows(client, company="PT DUP Tbk")["items"][0]["cells"]["net_income"]
        assert "disputed" not in cell
        assert cell["normalized_value"] == 17_527_130_084.0
        assert wrong.id != cell["id"]


# --------------------------------------------------------------------------
# the Excel export
#
# The download is the grid in a different shape, so the risk is not that it is
# wrong but that it is a *different* answer: a page-limited file, a filter
# quietly dropped, a missing figure written as zero. Each test pins one of those.

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _workbook(response) -> Workbook:
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MIME
    # The browser has to save this rather than try to render it.
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment;")
    assert ".xlsx" in disposition
    return load_workbook(io.BytesIO(response.content))


def _export(client, **params):
    return client.get("/api/results/summary/export", params=params)


def _sheet_rows(ws):
    return list(ws.iter_rows(min_row=2, values_only=True))


class TestExcelExport:
    def test_it_is_a_workbook_with_the_grid_as_its_first_sheet(self, client):
        wb = _workbook(_export(client))
        assert wb.sheetnames[0] == "Summary"
        head = [c.value for c in wb["Summary"][1]]
        assert head[:5] == [
            "Company", "Year", "Currency", "Sub-sector", "Pencatatan",
        ]
        assert head[5:7] == ["Total assets", "Total equity"]
        assert head[-2:] == ["Checks failed", "Source documents"]

    def test_figures_are_numbers_so_the_recipient_can_total_them(self, client):
        """Text figures cannot be summed, which is the main reason to ask for a
        workbook rather than a picture of the grid."""
        row = next(r for r in _sheet_rows(_workbook(_export(client))["Summary"])
                   if r[0] == "PT AAA Tbk")
        # Column 4 is the listing date; the first figure follows it.
        assert row[5] == 1_000_000.0
        assert isinstance(row[5], (int, float))

    def test_an_absent_figure_stays_empty_rather_than_becoming_zero(self, client):
        """A blank means "no report disclosed this line", which is a different
        claim from nil. A 0 in that cell would be a figure nobody reported."""
        row = next(r for r in _sheet_rows(_workbook(_export(client))["Summary"])
                   if r[0] == "PT CCC Tbk")
        assert row[0] == "PT CCC Tbk"
        assert row[5] is None, "CCC has sales but no total_assets"

    def test_it_covers_every_row_not_just_the_visible_page(self, client, tmp_path):
        """The grid pages at 50. An export that inherited the page size would
        silently drop everything past it, which is the worst way for a download
        to be wrong: it looks complete."""
        Session = get_session_factory(
            get_engine(client.app.dependency_overrides[get_config]())
        )
        with Session() as s:
            repo = Repository(s)
            doc, _ = repo.upsert_document(
                str(tmp_path / "many.pdf"), "hash-many", "PT MANY Tbk", "many.pdf"
            )
            repo.replace_values(doc, [
                # Distinct years, so these are 60 rows rather than 60 readings of
                # one figure.
                _value(repo, doc, "total_assets", 1_000.0 + i, year=1980 + i)
                for i in range(60)
            ])

        grid = _rows(client, page_size=50)
        assert grid["pagination"]["total"] == 65
        assert len(grid["items"]) == 50, "the fixture alone must exceed one page"

        ws = _workbook(_export(client))["Summary"]
        assert ws.max_row == 66  # header + every row

    def test_it_honours_the_filters_the_screen_is_showing(self, client):
        """Otherwise the file and the table answer two different questions."""
        companies = {
            r[0] for r in _sheet_rows(
                _workbook(_export(client, company="PT BBB"))["Summary"]
            )
        }
        assert companies == {"PT BBB Tbk"}

        years = [
            r[1] for r in _sheet_rows(_workbook(_export(client, year=2023))["Summary"])
        ]
        assert years == [2023]

    def test_the_currency_filter_narrows_the_file_too(self, client):
        rows = _sheet_rows(_workbook(_export(client, currency="USD"))["Summary"])
        assert [r[0] for r in rows] == ["PT BBB Tbk"]
        # total_assets on that row is undetected, so USD filtering must leave it
        # empty rather than reporting it under the wrong currency.
        assert rows[0][5] is None
        assert rows[0][9] == 120_000.0

    def test_it_carries_the_same_columns_as_the_grid(self, client):
        """A file whose columns differ from the screen cannot be read against it."""
        grid = _rows(client)
        head = [c.value for c in _workbook(_export(client))["Summary"][1]]
        assert head[5:-2] == [grid["labels"][f] for f in grid["fields"]]

    def test_a_row_whose_figures_do_not_reconcile_says_so_in_the_file(self, client):
        """The grid badges this row. An export that dropped the warning would be
        more dangerous than the screen, because it travels further."""
        row = next(r for r in _sheet_rows(_workbook(_export(client))["Summary"])
                   if r[0] == "PT AAA Tbk")
        assert "assets_equals_liabilities_plus_equity" in row[-2]

    def test_it_records_the_filters_it_was_taken_under(self, client):
        """A file with no note of what it was narrowed to cannot be checked
        against the screen it came from."""
        notes = _workbook(_export(client, company="PT BBB"))["Notes"]
        text = "\n".join(str(c.value) for r in notes.iter_rows() for c in r if c.value)
        assert "PT BBB" in text
        assert "Blank figure" in text, "the caveats must travel with the data"

    def test_it_records_every_selected_subsector(self, client):
        """The sub-sectors travel as repeated parameters; the note must
        carry each one, not a Python list repr."""
        wb = _workbook(
            _export(client, subsector=["D2. Food & Beverage", "11. Crops"])
        )
        assert [r[0] for r in _sheet_rows(wb["Summary"])] == [
            "PT AAA Tbk",
            "PT AAA Tbk",
            "PT CCC Tbk",
        ]
        text = "\n".join(
            str(c.value) for r in wb["Notes"].iter_rows() for c in r if c.value
        )
        assert "D2. Food & Beverage" in text
        assert "11. Crops" in text
        assert "['" not in text, "the filter note must be readable, not a repr"

    def test_an_empty_result_is_still_a_valid_workbook(self, client):
        """No rows must not mean a 500: a download that fails on an empty filter
        is indistinguishable from a broken button."""
        ws = _workbook(_export(client, company="PT NOPE"))["Summary"]
        assert [c.value for c in ws[1]][:3] == ["Company", "Year", "Currency"]
        assert ws.max_row == 1


# --------------------------------------------------------------------------
# GET /api/summary/financials -- the corpus-level aggregates the home
# dashboard leads with. The contract: the same winner rules as the grid,
# the latest reported year per company, and coverage that counts
# companies, not readings.
# --------------------------------------------------------------------------


class TestFinancialPulse:
    def test_leaders_use_the_latest_year_not_the_largest_reading(self, client):
        """AAA's 2023 assets (900k) are smaller than its 2024 ones. A
        ranking keyed on the biggest single reading would show the stale
        year, so the latest reported year must win."""
        data = client.get("/api/summary/financials").json()
        assert [
            (r["company"], r["year"], r["value"])
            for r in data["leaders"]["total_assets"]
        ] == [
            ("PT BBB Tbk", 2024, 2_000_000.0),
            ("PT AAA Tbk", 2024, 1_000_000.0),
        ]

    def test_leaders_rank_per_field(self, client):
        data = client.get("/api/summary/financials").json()
        assert data["leaders"]["sales_and_revenue"] == [
            {"company": "PT AAA Tbk", "year": 2024, "value": 800_000.0}
        ]
        assert data["leaders"]["total_equity"] == [
            {"company": "PT AAA Tbk", "year": 2024, "value": 600_000.0}
        ]
        # Nobody reports the headline profit figure, so the ranking is
        # empty rather than full of zeros.
        assert data["leaders"]["total_profit_loss"] == []

    def test_subsectors_count_companies_once(self, client):
        data = client.get("/api/summary/financials").json()
        assert data["subsectors"] == [
            {"subsector": "11. Crops", "companies": 1},
            {"subsector": "D2. Food & Beverage", "companies": 1},
        ]
        # BBB declares no classification at all, and that is a state the
        # dashboard can say out loud.
        assert data["undeclared_subsectors"] == 1

    def test_coverage_counts_distinct_companies(self, client):
        data = client.get("/api/summary/financials").json()
        coverage = data["coverage"]
        # AAA and BBB both report assets; CCC's "sales" is a different
        # field from "sales_and_revenue" and must not count as it.
        assert coverage["total_assets"] == 2
        assert coverage["sales_and_revenue"] == 1
        # BBB reports net_income twice (dated and undated) -- one company.
        assert coverage["net_income"] == 1
        assert coverage["non_controlling_interest"] == 0

    def test_an_empty_corpus_pulses_empty(self, tmp_path, monkeypatch):
        monkeypatch.delenv("AI_PROVIDER", raising=False)
        cfg = load_config()
        cfg.database_url = f"sqlite:///{(tmp_path / 'db' / 'processing.db').as_posix()}"
        set_config(cfg)
        app = create_app(cfg)
        app.dependency_overrides[get_config] = lambda: cfg
        with TestClient(app) as c:
            data = c.get("/api/summary/financials").json()
        assert data["subsectors"] == []
        assert data["undeclared_subsectors"] == 0
        assert data["leaders"]["total_assets"] == []
        assert all(n == 0 for n in data["coverage"].values())
